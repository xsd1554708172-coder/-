# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('output')
OUT.mkdir(exist_ok=True)


def set_cell_text(cell, text, size=9.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = '宋体'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_shading(cell, fill='E7E6E6'):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('第 ')
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = 'PAGE'
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    paragraph.add_run(' 页')
    for r in paragraph.runs:
        r.font.name = '宋体'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        r.font.size = Pt(9)


def setup_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.3)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)
    sec.header_distance = Cm(1.2)
    sec.footer_distance = Cm(1.2)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = '宋体'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    if 'LegalTitle' not in styles:
        s = styles.add_style('LegalTitle', WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = '黑体'
        s._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        s.font.size = Pt(18)
        s.font.bold = True
        s.paragraph_format.space_after = Pt(8)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if 'LegalHeading1' not in styles:
        s = styles.add_style('LegalHeading1', WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = '黑体'
        s._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        s.font.size = Pt(14)
        s.font.bold = True
        s.paragraph_format.space_before = Pt(8)
        s.paragraph_format.space_after = Pt(4)

    if 'LegalHeading2' not in styles:
        s = styles.add_style('LegalHeading2', WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = '楷体'
        s._element.rPr.rFonts.set(qn('w:eastAsia'), '楷体')
        s.font.size = Pt(12)
        s.font.bold = True
        s.paragraph_format.space_before = Pt(5)
        s.paragraph_format.space_after = Pt(2)

    for section in doc.sections:
        add_page_number(section.footer.paragraphs[0])


def add_title(doc, title, subtitle=None):
    p = doc.add_paragraph(style='LegalTitle')
    p.add_run(title)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run(subtitle)
        r.font.name = '宋体'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        r.font.size = Pt(11)


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style='LegalHeading1' if level == 1 else 'LegalHeading2')


def add_body(doc, text='', first_line=True, bold_prefix=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(24)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    for r in p.runs:
        r.font.name = '宋体'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        r.font.size = Pt(12)
    return p


def add_numbered(doc, number, text, indent=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Pt(indent)
    p.paragraph_format.first_line_indent = Pt(-18)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(f'{number} ')
    r1.bold = True
    p.add_run(text)
    for r in p.runs:
        r.font.name = '宋体'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        r.font.size = Pt(12)
    return p


def add_info_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, (k, v) in enumerate(rows):
        set_cell_text(t.cell(i, 0), k, 10.5, True)
        set_cell_text(t.cell(i, 1), v, 10.5, False, WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_shading(t.cell(i, 0), 'E7E6E6')
        set_cell_margins(t.cell(i, 0)); set_cell_margins(t.cell(i, 1))
    return t


def add_signature_table(doc, start_no, count):
    headers = ['序号', '姓名', '身份证后4位', '个人情况\nA/B/C/D/E', '本人签名', '本人指印', '备注']
    t = doc.add_table(rows=count + 1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [0.8, 2.0, 1.8, 2.2, 2.2, 2.0, 3.1]
    for j, h in enumerate(headers):
        c = t.cell(0, j)
        set_cell_text(c, h, 8.5, True)
        set_cell_shading(c, 'D9E2F3')
        set_cell_margins(c)
        c.width = Cm(widths[j])
    set_repeat_table_header(t.rows[0])
    for i in range(1, count + 1):
        set_cell_text(t.cell(i, 0), start_no + i - 1, 9)
        for j in range(1, len(headers)):
            set_cell_text(t.cell(i, j), '\n', 9)
            set_cell_margins(t.cell(i, j), top=100, bottom=100)
        t.rows[i].height = Cm(0.85)
    return t


def add_signoff(doc, include_seal=False):
    t = doc.add_table(rows=3 if include_seal else 2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    items = [
        ('主持人签名：________________', '记录人签名：________________'),
        ('屯长签名：__________________', '见证人签名：________________'),
    ]
    for i, pair in enumerate(items):
        for j, txt in enumerate(pair):
            set_cell_text(t.cell(i, j), txt, 11, False, WD_ALIGN_PARAGRAPH.LEFT)
            t.cell(i, j)._tc.get_or_add_tcPr().remove_all(qn('w:tcBorders')) if False else None
    if include_seal:
        set_cell_text(t.cell(2, 0), '山铁村民小组盖章处：\n\n\n', 11, False, WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(t.cell(2, 1), '日期：2026年____月____日', 11, False, WD_ALIGN_PARAGRAPH.LEFT)
    return t


# ---------------- Document 1 ----------------
def build_meeting_doc():
    doc = Document()
    setup_document(doc)
    add_title(doc, '山铁屯确认合同有效纠纷应诉会议材料', '屯长宣读稿、会议重点、会议决议及签名按印表')
    add_info_table(doc, [
        ('案号', '（2026）桂1424民初1755号'),
        ('案件', '黄全和诉大新县堪圩乡芦山村山铁村民小组确认合同有效纠纷'),
        ('会议时间', '2026年____月____日____时'),
        ('会议地点', '________________________________'),
        ('主持人／记录人', '主持人：____________　记录人：____________'),
    ])

    add_heading(doc, '第一部分　屯长宣读稿')
    add_body(doc, '各位乡亲：黄全和已经向大新县人民法院起诉，请求确认2001年《土地承包合同书》和2012年《〈土地承包合同书〉补充协议》有效，并要求继续履行。法院目前只是受理案件，尚未判定合同有效。我们今天开会，是为了把全屯对本案的真实意见、事实线索和可以接受的处理方案说明清楚。')
    add_body(doc, '大家必须根据本人经历如实表达。确实同意过使用土地的，要说明当时同意了什么；确实签过字或者按过指印的，要如实承认；没有见过完整合同、不知道30年期限或者具体价款的，也要分别说明。任何人不得要求大家统一说法。记不清的事项就写“记不清”，无法辨认签名或指印的就写“无法确认，请求依法鉴定”。')

    add_heading(doc, '一、为什么经过多年才集中提出异议')
    add_numbered(doc, '1.', '全屯并不否认黄全和长期实际种植、使用案涉土地。但“知道有人种地”与“知道并同意一份从2001年持续至2031年、每亩每年32元、每十年付款一次的完整书面合同”属于不同事实。')
    add_numbered(doc, '2.', '多数屯民长期没有保管2001年完整合同，也没有被集中告知原告现在主张的30年期限、具体价款及付款安排。2026年发生收地争议并在调解、诉讼中集中查看相关材料后，屯民才逐项核对合同正文、2001年签字按印单、2012年补充协议、收条及账单。')
    add_numbered(doc, '3.', '核对后出现多种情况：有人只同意过把地交给黄全和种植，却不知道期限为30年；有人认为签名并非本人所写；有人表示只签过名、没有按过现有指印；有人虽可能签名，但没有看过含期限和价款的完整正文。发现这些问题后，村民小组随即提出真实性、授权和效力异议。')
    add_numbered(doc, '4.', '因此，本屯现在提出异议的重点并非多年后单纯嫌价格偏低，而是要求法院查明：当年是否真正讨论、说明并同意了30年期限及全部主要条件；签字、指印是否真实；签约人员是否获得授权；2012年三名签署人员是否有权代表全屯确认至2031年的剩余承包期。')

    add_heading(doc, '二、本案需要法院重点查明的事实')
    add_numbered(doc, '1.', '2001年1月6日手写材料上究竟讨论了哪些事项。该页现有复印件主要写有同意把“陇梅”发包给黄全和，但页面上未见30年期限、每亩32元、付款周期等完整条款。需要查明参会人员当时听到的内容以及是否授权负责人另行确定期限、价格。')
    add_numbered(doc, '2.', '2001年正式合同及签字按印单是否有原件，甲方签署人员当时担任何种职务、权限从何而来；乡政府、芦山村委会是否保存合同所称的一式四份存档文本。')
    add_numbered(doc, '3.', '2012年补充协议由李培光、农云伟、李飞球等人签署时，是否召开村民会议、是否具有明确授权、是否向全屯公开“剩余期限至2031年”和一次收取25600元等事项。')
    add_numbered(doc, '4.', '即使25600元确已收取且部分用于铺路，也要分别查明款项来源、完整账目、是否公开以及全屯是否知道该笔钱对应的是2011年至2031年的承包费。看见道路修建或使用道路，本身不能直接代替对补充协议完整内容的知情和同意。')
    add_numbered(doc, '5.', '对存在具体异议的签名、指印及印章，应当在原件具备检验条件时依法核验或鉴定。')

    add_heading(doc, '三、山铁屯在本案中的主要主张和诉求')
    add_numbered(doc, '1.', '不同意原告关于确认2001年合同、2012年补充协议有效并继续履行的诉讼请求，请求人民法院依法驳回。')
    add_numbered(doc, '2.', '请求原告提交2001年村民表决签字单、正式合同、2012年补充协议、收条、账簿及相关附件原件，并调取乡政府、芦山村委会保存的合同和备案材料。')
    add_numbered(doc, '3.', '请求法院审查签署人员的代表权限和授权范围。加盖印章、个别人收款或修路支出，不能当然证明全屯已知悉并追认至2031年的全部合同内容。')
    add_numbered(doc, '4.', '对具体提出异议的签名、指印和印章，申请依法进行笔迹、指印或者印章鉴定。')
    add_numbered(doc, '5.', '如法院认为返还土地、确认合同不发生效力或者无效、清算已付款和土地使用费等主动请求需要通过反诉或者另案处理，同意由村民小组依法另行办理。')
    add_numbered(doc, '6.', '在案件审理期间，村民不得自行损毁作物、设备或者扩大冲突。有关土地现状和生产安排，服从法院依法处理或者双方书面和解。')

    add_heading(doc, '四、屯民需要协助准备的证据')
    add_numbered(doc, '1.', '逐人查看2001年1月6日签字按印单，写明本人签名、指印的真实意见，以及当时听到的期限、价格和授权内容。')
    add_numbered(doc, '2.', '提供2000年至2012年前后自然形成的签名材料，如户籍、银行、土地、婚姻登记、收据等，供法院决定是否鉴定。')
    add_numbered(doc, '3.', '2001年及2012年的经办人分别说明会议、签署、收款、记账、铺路和村务公开情况，并保留原始账本、票据。')
    add_numbered(doc, '4.', '说明2026年4月会议为什么理解为2021年到期，以及本人何时第一次完整看到原告据以主张至2031年的合同材料。没有书面依据的，不得补造材料。')

    add_heading(doc, '五、可以提交法院讨论的和解方向')
    add_numbered(doc, '方案一：', '原告撤回确认旧合同有效并继续履行的请求；约定现有作物合理收获后返还土地；双方按照真实支付、合理土地使用价值和有凭证的必要投入进行结算、抵扣。')
    add_numbered(doc, '方案二：', '如全屯依法重新开会并同意继续由原告使用土地，双方重新确定面积、期限、价格、付款账户、调价办法、违约责任和土地恢复义务，另行签订新合同；旧合同不再作为继续占地的依据。')
    add_body(doc, '以上和解意见只用于法院主持下解决纠纷，不表示本屯承认旧合同有效。最终期限、金额及权利放弃须再次经有效集体决定，屯长和一般诉讼代表无权单独作出重大让步。', first_line=True)

    doc.add_page_break()
    add_heading(doc, '第二部分　山铁屯村民小组会议决议')
    add_body(doc, '2026年____月____日，山铁屯就（2026）桂1424民初1755号确认合同有效纠纷召开会议。主持人已将本文件全文宣读，并对原告的诉讼请求、主要证据、本屯可能面临的不利事实及应诉意见进行了说明。参会人员根据本人真实意思表决，形成如下决议：')
    add_numbered(doc, '一、', '同意以本文件所列事实意见和诉讼立场进行应诉，请求法院驳回原告要求确认2001年合同及2012年补充协议有效并继续履行的诉讼请求。')
    add_numbered(doc, '二、', '同意申请原告提交全部原件，申请法院调查调取乡政府、芦山村委会存档材料；对有具体异议且具备条件的签名、指印、印章申请依法鉴定。')
    add_numbered(doc, '三、', '同意由屯长________________及推选代表________________、________________负责向法院提交答辩、质证意见、证据和一般性调解意见，接收并转达法院通知。')
    add_numbered(doc, '四、', '上述人员无权单独承认旧合同有效、放弃主要抗辩、同意继续承包期限、确定最终和解金额或者处分土地。涉及上述事项，须另行召开会议形成决议。')
    add_numbered(doc, '五、', '每名签署人只对本人知道的事实负责。签名按印表中A至E类型可以多选；选择C项的人员，另行填写个人真实性异议说明并配合法院核验。')
    add_numbered(doc, '六、', '会议期间及诉讼期间，任何人不得诱导他人统一陈述，不得擅自毁损原告现有作物、设备或以冲突方式处理争议。')

    add_info_table(doc, [
        ('山铁屯总户数', '______户'),
        ('本次应通知人数／户数', '______人／户'),
        ('实际参会人数／户数', '______人／户'),
        ('同意本决议', '______人／户'),
        ('不同意／弃权', '不同意______人／户；弃权______人／户'),
    ])
    add_signoff(doc, include_seal=True)

    doc.add_page_break()
    add_heading(doc, '第三部分　签名按印说明')
    add_body(doc, '本人已听取或者阅读本文件全文，理解原告的诉讼请求、本屯的答辩主张、证据准备事项及和解范围，同意／不同意在下表如实填写。姓名以身份证或者户口簿为准。本人亲自签名并按右手拇指印；个人无须盖章，村民小组在会议决议处统一盖章。')
    add_body(doc, '个人情况代码可多选：A＝当时未见完整合同，不知道30年期限或者具体价款；B＝当年曾原则同意土地使用，但没有签署或者没有同意现存完整合同；C＝对现存签名、指印真实性有异议，申请依法核验或鉴定；D＝签名可能系本人所写，但本人未按现有指印，或者签名时不知道完整条件；E＝其他情况，在备注栏写明。')
    add_signature_table(doc, 1, 20)
    add_body(doc, '本页签署人数：______人。核对人签名：________________。', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_page_break()
    add_heading(doc, '山铁屯集体应诉意见签名按印表（续页）')
    add_signature_table(doc, 21, 20)
    add_body(doc, '本页签署人数：______人；两页合计：______人。', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_body(doc, '核对人签名：________________　屯长签名：________________　日期：2026年____月____日', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)

    path = OUT / '山铁屯会议宣读_会议重点与签名按印材料.docx'
    doc.save(path)
    return path


# ---------------- Document 2 ----------------
def build_litigation_doc():
    doc = Document()
    setup_document(doc)
    add_title(doc, '（2026）桂1424民初1755号', '民事答辩及相关申请材料')
    add_info_table(doc, [
        ('提交法院', '大新县人民法院（雷平人民法庭）'),
        ('原告', '黄全和'),
        ('被告', '大新县堪圩乡芦山村山铁村民小组'),
        ('案由', '确认合同有效纠纷'),
        ('开庭时间', '2026年10月21日09时30分'),
        ('材料日期', '2026年____月____日'),
    ])
    add_heading(doc, '材料目录')
    for n, item in enumerate([
        '民事答辩状',
        '对原告证据的书面质证意见',
        '责令提交书证原件及调查收集证据申请书',
        '司法鉴定申请书（待填写具体人员和位置）',
        '被告证据目录',
        '调解意见书',
        '提交材料清单',
    ], 1):
        add_numbered(doc, f'{n}.', item)

    doc.add_page_break()
    add_title(doc, '民事答辩状')
    add_body(doc, '答辩人（被告）：大新县堪圩乡芦山村山铁村民小组。', first_line=False)
    add_body(doc, '负责人：黄迎春，组长。', first_line=False)
    add_body(doc, '被答辩人（原告）：黄全和。', first_line=False)
    add_body(doc, '案号：（2026）桂1424民初1755号。', first_line=False)

    add_heading(doc, '答辩请求')
    add_numbered(doc, '一、', '依法驳回原告关于确认2001年1月7日《土地承包合同书》及2012年11月15日《〈土地承包合同书〉补充协议》有效并继续履行的全部诉讼请求。')
    add_numbered(doc, '二、', '本案诉讼费用由原告承担。')
    add_numbered(doc, '三、', '责令原告提交其所援引的村民表决签字单、2001年合同、2012年补充协议、收条、土地承包费使用账单及相关附件原件；对存在具体异议的签名、指印、印章依法核验或者鉴定。')
    add_numbered(doc, '四、', '调查调取堪圩乡人民政府、芦山村民委员会依法保存的合同、备案、会议及村务档案材料。')

    add_heading(doc, '事实与理由')
    add_heading(doc, '一、原告应当先证明两份合同真实成立、签署人员具有权限，不能仅以复印件、印章和长期使用土地直接推定合同有效', level=2)
    add_body(doc, '原告要求法院确认合同有效并继续履行，应当对合同原件、签名指印真实性、签署人员权限、集体决策过程以及其已按约履行承担举证责任。原告证据目录明确记载证据1至证据9均为复印件。现有复印件尚无法完整核对纸张、墨迹、签名、指印、印章及附件之间的对应关系。被告请求原告出示原件，并请求法院按照《最高人民法院关于民事诉讼证据的若干规定》关于原件、私文书真实性、复制件证明力和控制证据一方拒不提交之法律后果的规定进行审查。')
    add_body(doc, '2001年合同正文还写明合同一式四份，甲乙双方、堪圩乡政府、芦山村委会各执一份；2012年补充协议也写明交芦山村委会存档一份。堪圩乡政府2026年出具的说明却称未找到山铁经联社的档案资料。该说明本身不能直接证明合同无效，也不能证明合同真实有效，但足以表明存档链条需要进一步核实。')

    add_heading(doc, '二、2001年1月6日手写同意发包材料未载明30年期限、每亩32元及付款方式，不能未经审查就推定名单上的所有人员同意正式合同全部条款', level=2)
    add_body(doc, '从原告提交的复印件可以看出，2001年1月6日手写材料主要记载同意将“陇梅”发包给本屯村民黄全和，并列有多人姓名和指印。该页未见2001年至2031年的30年期限、每亩每年32元、每十年支付12800元等关键条款。上述期限、价格和付款安排出现在次日的正式合同正文。')
    add_body(doc, '即使部分村民当年原则上同意由原告使用、种植案涉土地，也仍应查明会议是否讨论过完整期限、价格和付款条件，是否授权特定人员另行确定这些事项。原则同意发包与接受现存完整合同的全部主要条款不能当然等同。')
    add_body(doc, '部分在世人员已经对该手写材料中本人名下的签名、指印提出具体异议；另有人员表示当年只知道有人承包，并不知道期限长达30年。被告将提交个人说明及历史笔迹线索，请求在原件具备检验条件时依法处理鉴定申请。')

    add_heading(doc, '三、2001年正式合同的甲方签署权限、印章和备案情况尚未查清，30年期限在抽象上可能被允许，并不能反向证明签约程序和授权当然有效', level=2)
    add_body(doc, '原告起诉状所称面积40亩、期限2001年1月7日至2031年1月7日、承包费每亩每年32元、总额38400元，均来自其提交的正式合同复印件。被告对原件、甲方签署人员身份及代表权限、村民会议决策内容、印章形成与使用过程提出异议。')
    add_body(doc, '原告引用耕地承包期30年的规定，只能说明30年期限并非当然受禁止。案涉安排究竟属于家庭承包还是以支付承包费方式进行的其他承包，仍需结合土地性质、发包方式和集体决策情况判断。期限可能合法，不代表合同必然经过真实表决、充分告知和合法授权。')
    add_body(doc, '2001年签约时应依当时有效的法律审查。依当时《合同法》第四十八条、第五十条，无代理权、超越代理权订立合同未经追认的，对被代理人不发生效力；负责人越权订立合同时，还需审查相对人是否知道或者应当知道其超越权限。原告系本屯成员，对本屯组织结构和重大集体财产事项的决策方式具有了解条件。')

    add_heading(doc, '四、2012年补充协议由少数人员签署，原告尚需证明三名签署人员取得了确认剩余承包期至2031年并一次收取25600元的集体授权', level=2)
    add_body(doc, '2012年补充协议载明李培光、农云伟、李飞球等人为甲方代表，并将2011年至2031年的剩余承包费25600元改为一次性支付。原告目前没有提交2012年村民会议记录、表决名册、授权书或者村务公开材料，用以证明三名人员有权代表全屯作出上述重大决定。')
    add_body(doc, '公章和个别人员签字是法院审查的重要事实，但公章不能替代对代表权限、授权范围以及原告是否合理核查权限的审查。2012年补充协议也不能在缺乏授权证明的情况下当然补正2001年合同存在的知情、授权和真实性争议。')

    add_heading(doc, '五、收取25600元和所谓用于铺路的事实，应与“全屯知悉并追认承包至2031年”分别审查', level=2)
    add_body(doc, '被告不回避原告提交了25600元收条及两页手写支出账单。相关经办人员是否真实收款、款项如何使用，应当以原件、完整账簿、票据、经办人陈述和村务公开材料查明。现有两页账单分别出现6708元和19216元等数字，合计与25600元并不完全一致，且缺少完整年份、资金来源和原始凭证。')
    add_body(doc, '即使法院最终认定25600元确已支付且部分用于铺路，也只能证明有关人员收取、支出过款项。要认定全屯事后追认补充协议，仍需证明有权决策主体知道该款项对应的是至2031年的剩余20年承包费，知道补充协议完整内容，并在此基础上作出明确认可或者足以表明认可的行为。村民看见道路修建、使用道路，不能在缺乏上述知情证据时直接推定其接受全部合同内容。')
    add_body(doc, '补充协议记载剩余承包期为2011年1月7日至2031年1月7日，收条复印件记载期间为2011年12月30日至2031年12月30日。两者起止日期不一致，亦需由原件和经办人员解释。')

    add_heading(doc, '六、被告经过多年才集中提出异议，有具体事实原因；长期使用土地属于审查因素，但不能代替对真实知情和授权的证明', level=2)
    add_body(doc, '被告承认原告长期实际种植、使用案涉土地。多数普通屯民长期知道有人种地，但没有保管完整合同，也未被集中告知原告现在主张的30年期限、价格和付款安排。2026年发生收地争议，并在乡政府调解、人民法院诉讼中集中查看原告材料后，屯民才逐项核对签名、指印和合同条款，并陆续提出异议。')
    add_body(doc, '2026年4月20日会议材料写有“2001年至2021年到期”，反映当时参会村民对期限的实际认识。其形成依据还需由会议组织者和知情人员说明。该材料至少与原告所称全屯25年来一直明知并认可至2031年的证明目的存在矛盾。')
    add_body(doc, '被告请求法院在维护交易稳定的同时，结合当地村民在2001年、2012年的文化程度、语言表达、会议方式和信息公开条件，具体查明是否向成员说明过30年期限、价款及一次性付款所代表的法律后果。该背景用于判断真实知情和意思表示，不以人数、情绪替代法律和证据。')

    add_heading(doc, '七、现有事实不足以支持原告要求两份合同当然有效并继续履行', level=2)
    add_body(doc, '本案存在原件缺失或尚未出示、表决材料未载关键条款、签名指印具体异议、代表权限不明、2012年授权材料缺失、收款与追认之间证明链条不完整等问题。原告关于合同程序完全合法、全屯长期无异议并已追认的主张，尚未达到足以排除上述合理疑问的程度。')
    add_body(doc, '综上，请求人民法院依法查明事实，驳回原告全部诉讼请求。')

    add_body(doc, '此致', first_line=False)
    add_body(doc, '大新县人民法院', first_line=False)
    add_body(doc, '答辩人：大新县堪圩乡芦山村山铁村民小组（盖章）', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_body(doc, '负责人签名：________________', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_body(doc, '2026年____月____日', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_page_break()
    add_title(doc, '对原告证据的书面质证意见')
    add_body(doc, '被告依据原告证据目录及现有复印件，提出以下初步质证意见。最终意见以原告当庭出示原件、播放完整视频及法院核对情况为准。')
    headers = ['序号', '原告证据', '真实性／形式意见', '证明目的意见']
    rows = [
        ('1', '堪圩乡政府《情况说明》', '对加盖政府印章的文件形式真实性以原件核对为准。', '该说明仅称未找到山铁经联社档案，并建议村民小组处理后续事项；不能证明2001年、2012年文件真实、授权充分或者合同有效。'),
        ('2', '2001年1月6日村民表决签字单', '系复印件；部分人员对本人名下签名、指印提出具体异议，请求出示原件并核验。', '该页可用于证明当时曾讨论发包事项，但页面未载30年期限、每亩32元及付款周期，不能单独证明所有人同意正式合同全部条款。'),
        ('3', '2001年《土地承包合同书》', '系复印件；要求提交甲乙双方、乡政府、村委会各执文本原件，核对签署、印章和附件。', '30年期限写入合同，不等于村民会议已讨论并授权；甲方代表权限、备案、履行情况仍需证明。'),
        ('4', '2012年补充协议', '系复印件；对三名签署人员权限、公章使用和形成过程有异议，要求原件。', '原告未提交2012年集体决议和授权材料，不能仅凭三人签字盖章证明全屯同意至2031年。'),
        ('5', '25600元收条', '系复印件；需核对原件、收款人员、日期和款项交付情况。收条期间与补充协议期间不一致。', '即使证明付款，也不能单独证明全屯知悉并追认补充协议全部内容。'),
        ('6', '土地承包费使用账单', '两页手写复印件，无完整账簿、发票、年份和村务公开材料，金额之间需要核对。', '至多证明部分人员记载过铺路支出，不能单独证明资金来源、全部支出真实或全屯追认合同。'),
        ('7', '2026年4月20日收回土地签字单', '如有原件，被告对会议召开和村民签字的基本事实予以确认，具体内容以原件为准。', '该证据反映参会村民当时理解合同至2021年，并明确反对继续承包；不能据此反推2001年、2012年合同当然有效。'),
        ('8', '现场照片', '对拍摄时间、地点、人员和完整背景需要核实。', '照片可反映现场存在竹桩、线绳等，不能单独证明由谁设置、是否造成损害，更不能证明合同效力。'),
        ('9', '乡政府调解申请书', '系原告单方制作并提交的申请材料。', '只能证明原告曾提出调解请求和单方陈述，不能证明其中合同效力、侵权或50万元损失主张成立。'),
        ('10', '视频光盘', '要求当庭完整播放、提供复制件并说明拍摄人、时间、地点和原始载体。播放前暂不作最终意见。', '村民在土地现场讨论不能直接证明存在侵权，也不能证明两份合同有效。'),
    ]
    t = doc.add_table(rows=1, cols=4)
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        set_cell_text(t.cell(0,j), h, 9, True); set_cell_shading(t.cell(0,j), 'D9E2F3'); set_cell_margins(t.cell(0,j))
    set_repeat_table_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            set_cell_text(cells[j], val, 8.5, False, WD_ALIGN_PARAGRAPH.LEFT if j else WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_margins(cells[j])

    add_body(doc, '质证人：大新县堪圩乡芦山村山铁村民小组（盖章）', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_body(doc, '负责人／代理人：________________　2026年____月____日', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_page_break()
    add_title(doc, '责令提交书证原件及调查收集证据申请书')
    add_body(doc, '申请人：大新县堪圩乡芦山村山铁村民小组。', first_line=False)
    add_body(doc, '案号：（2026）桂1424民初1755号。', first_line=False)
    add_heading(doc, '申请事项')
    items = [
        '责令原告提交2001年1月6日村民表决签字单原件。',
        '责令原告提交2001年1月7日《土地承包合同书》原件及其附件。',
        '责令原告提交2012年11月15日补充协议原件、25600元收条原件。',
        '责令掌握账目的人员提交与25600元有关的完整原始账簿、票据、付款凭证、村务公开记录。',
        '向堪圩乡人民政府、芦山村民委员会调查调取2001年合同所称备案、存档文本和有关盖章记录。',
        '向芦山村民委员会调查调取2012年补充协议所称存档文本，以及2012年前后山铁屯负责人、印章保管和会议记录材料。',
    ]
    for i, txt in enumerate(items, 1): add_numbered(doc, f'{i}.', txt)
    add_heading(doc, '事实与理由')
    add_body(doc, '原告以相关私文书证明合同有效，但其证据目录明确载明证据1至9均为复印件。相关原件直接关系签名、指印、印章、签署权限、文件形成和履行事实。被告无法自行取得原告、乡政府、村委会及历史经办人员控制的上述材料，故申请法院责令提交或者依法调查收集。')
    add_body(doc, '如控制书证的一方无正当理由拒不提交，请求法院依照民事证据规则，结合全案事实承担相应不利后果。')
    add_body(doc, '此致', first_line=False)
    add_body(doc, '大新县人民法院', first_line=False)
    add_body(doc, '申请人：大新县堪圩乡芦山村山铁村民小组（盖章）', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_body(doc, '负责人签名：________________　2026年____月____日', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_page_break()
    add_title(doc, '司法鉴定申请书')
    add_body(doc, '申请人：大新县堪圩乡芦山村山铁村民小组。', first_line=False)
    add_body(doc, '案号：（2026）桂1424民初1755号。', first_line=False)
    add_heading(doc, '申请事项')
    add_body(doc, '在原告提交相关原件且具备检验条件后，对下列具体签名、指印或者印章依法委托具备资质的鉴定机构进行鉴定：')
    at = doc.add_table(rows=6, cols=5)
    at.style = 'Table Grid'; at.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(['序号','文件及页码','争议姓名／印章','申请鉴定项目','比对样本线索']):
        set_cell_text(at.cell(0,j), h, 9, True); set_cell_shading(at.cell(0,j),'D9E2F3')
    for i in range(1,6):
        set_cell_text(at.cell(i,0), i, 9)
        for j in range(1,5): set_cell_text(at.cell(i,j), '\n\n', 9, False, WD_ALIGN_PARAGRAPH.LEFT)
    add_heading(doc, '事实与理由')
    add_body(doc, '原告提交的2001年表决签字单、正式合同、2012年补充协议及收条中存在签名、指印或者印章。部分相关人员经查看复印件后明确否认本人签名、指印，或者表示无法确认并愿意提供历史笔迹材料。上述真实性与合同是否真实成立、签署人员是否获得授权直接相关，申请依法鉴定。')
    add_body(doc, '申请人将在法院指定期限内补充具体争议位置、自然形成的历史样本，并按照法院要求配合办理。对无法鉴定或者样本不足的事项，请结合其他证据综合判断。')
    add_body(doc, '此致', first_line=False)
    add_body(doc, '大新县人民法院', first_line=False)
    add_body(doc, '申请人：大新县堪圩乡芦山村山铁村民小组（盖章）', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_body(doc, '负责人签名：________________　2026年____月____日', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_page_break()
    add_title(doc, '被告证据目录')
    headers = ['序号','证据名称','份数／页数','证明目的','原件情况']
    evidence_rows = [
        ('1','山铁屯本次应诉会议决议及签名按印表','','证明被告现行集体应诉立场、授权范围及调解范围','原件／复印件'),
        ('2','相关人员个人历史事实说明','','证明2001年、2012年的知情、签署、授权及追认情况','原件'),
        ('3','对争议签名、指印的个人异议说明','','证明具体异议及申请鉴定的必要性','原件'),
        ('4','相关人员历史笔迹材料','','作为法院审查是否鉴定及确定样本的线索','原件核验、复印件提交'),
        ('5','2026年4月20日会议材料','','证明参会村民当时对合同期限的真实认识及集体反对继续承包','原件／复印件'),
        ('6','乡政府、村委会调解或沟通记录','','证明纠纷发生、相关材料集中展示及被告提出异议的经过','待补充'),
        ('7','2012年收款、铺路完整账目及经办人说明','','证明款项来源、使用、公开以及是否形成知情追认','待调取／待补充'),
        ('8','其他：________________','','________________________________','________'),
    ]
    et = doc.add_table(rows=1, cols=5); et.style='Table Grid'; et.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(headers): set_cell_text(et.cell(0,j),h,9,True); set_cell_shading(et.cell(0,j),'D9E2F3')
    set_repeat_table_header(et.rows[0])
    for row in evidence_rows:
        cells=et.add_row().cells
        for j,val in enumerate(row): set_cell_text(cells[j],val,8.5,False,WD_ALIGN_PARAGRAPH.LEFT if j else WD_ALIGN_PARAGRAPH.CENTER)
    add_body(doc, '提交人：大新县堪圩乡芦山村山铁村民小组（盖章）', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_body(doc, '负责人／代理人：________________　2026年____月____日', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_page_break()
    add_title(doc, '调解意见书')
    add_body(doc, '被告在不承认2001年合同及2012年补充协议有效的前提下，为实质解决纠纷，提出以下可讨论方案：')
    add_heading(doc, '方案一　终止旧合同争议并返还土地', level=2)
    add_numbered(doc, '1.', '原告撤回确认旧合同有效并继续履行的诉讼请求；双方确认旧文件不再继续履行。')
    add_numbered(doc, '2.', '原告于________年____月____日前，或者本季现有作物合理收获后，将案涉土地交还被告。')
    add_numbered(doc, '3.', '双方依据证据核算原告已付款项、土地实际使用期间合理使用价值、有凭证的必要投入及现有作物处理费用，相互抵扣后一次结清。')
    add_numbered(doc, '4.', '结清后双方不再依据旧合同主张继续履行或违约责任；交地范围、现状和界址制作书面清单。')
    add_heading(doc, '方案二　依法重新协商并另签新合同', level=2)
    add_numbered(doc, '1.', '须经被告依法召开会议形成有效集体决定。')
    add_numbered(doc, '2.', '重新明确面积、四至、用途、期限、价格、付款账户、支付周期、价格调整、违约责任、地力保护和期满恢复义务。')
    add_numbered(doc, '3.', '新合同生效后，2001年合同及2012年补充协议不再作为继续占有、使用土地的依据。')
    add_body(doc, '本意见只作为法院主持调解的基础。屯长或诉讼代表无权未经新的集体决议，单独承认旧合同有效、确定继续使用期限、最终金额或者放弃主要权利。')
    add_body(doc, '被告：大新县堪圩乡芦山村山铁村民小组（盖章）', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_body(doc, '负责人签名：________________　2026年____月____日', first_line=False, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_page_break()
    add_title(doc, '提交材料清单')
    submission = [
        ('1','民事答辩状','正本1份、副本1份',''),
        ('2','书面质证意见','正本1份、副本1份',''),
        ('3','责令提交原件及调查取证申请书','1份',''),
        ('4','司法鉴定申请书','1份','具体人员和位置填写后提交'),
        ('5','被告证据目录及证据复印件','按法院要求','原件开庭携带'),
        ('6','应诉会议决议及签名按印表','原件1份、复印件1份',''),
        ('7','负责人身份证明书／授权委托书','各1份','如有代理人'),
        ('8','调解意见书','1份','可单独密封或调解时提交'),
    ]
    st=doc.add_table(rows=1,cols=4); st.style='Table Grid'; st.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,h in enumerate(['序号','材料名称','份数','备注']): set_cell_text(st.cell(0,j),h,9,True); set_cell_shading(st.cell(0,j),'D9E2F3')
    for row in submission:
        cells=st.add_row().cells
        for j,val in enumerate(row): set_cell_text(cells[j],val,9,False,WD_ALIGN_PARAGRAPH.LEFT if j else WD_ALIGN_PARAGRAPH.CENTER)
    add_body(doc, '提交人签名：________________　法院接收人：________________　提交日期：2026年____月____日', first_line=False)

    path = OUT / '（2026）桂1424民初1755号民事答辩及相关申请材料.docx'
    doc.save(path)
    return path


if __name__ == '__main__':
    p1 = build_meeting_doc()
    p2 = build_litigation_doc()
    print(p1)
    print(p2)
